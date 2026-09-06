from pathlib import Path

from openpyxl import Workbook

from historical_bloodlines.application.services.assembler import GenealogyAssembler
from historical_bloodlines.application.services.build_genealogy import BuildGenealogyUseCase
from historical_bloodlines.application.services.parser import GenealogyRowParser
from historical_bloodlines.infrastructure.excel import ExcelGenealogyReader
from historical_bloodlines.infrastructure.graph.renderer import GraphvizGenealogyRenderer


HEADERS = [
    "№",
    "Имя",
    "Титул",
    "Начало правления",
    "Конец правления",
    "Дети",
    "Брак",
]


def test_optional_note_is_carried_from_excel_to_domain(tmp_path: Path) -> None:
    source = tmp_path / "notes.xlsx"
    note = "* Примечание о правах Генриха Тюдора на английскую корону."

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Тюдоры"
    sheet.append([*HEADERS, "Примечание"])
    sheet.append([1, "Генрих VII Тюдор", None, 1485, 1509, None, None, note])
    workbook.save(source)

    raw = ExcelGenealogyReader().read(source)[0].rows[0]
    assert raw.note_raw == note

    parsed = GenealogyRowParser().parse(raw)
    assert parsed.note == note

    genealogy = GenealogyAssembler().assemble((parsed,))
    person = next(iter(genealogy.persons.values()))
    assert person.note == note


def test_note_is_rendered_as_table_footnote_and_marks_person(tmp_path: Path) -> None:
    source = tmp_path / "notes.xlsx"
    output = tmp_path / "notes.svg"
    note = "* Примечание Генриха Тюдора."

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Тюдоры"
    sheet.append([*HEADERS, "Примечание"])
    sheet.append([1, "Генрих VII Тюдор", None, 1485, 1509, None, None, note])
    workbook.save(source)

    rendered, warnings = BuildGenealogyUseCase().execute(source, output)

    assert warnings == ()
    svg = next(rendered.glob("*.svg")).read_text(encoding="utf-8")
    assert "Генрих VII Тюдор*" in svg
    assert "Примечание Генриха Тюдора." in svg


def test_renderer_adds_asterisk_to_footnote_when_excel_omits_it(tmp_path: Path) -> None:
    source = tmp_path / "notes.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Тюдоры"
    sheet.append([*HEADERS, "Примечание"])
    sheet.append([1, "Генрих VII Тюдор", None, None, None, None, None, "Текст сноски"])
    workbook.save(source)

    raw = ExcelGenealogyReader().read(source)[0].rows[0]
    genealogy = GenealogyAssembler().assemble((GenealogyRowParser().parse(raw),))
    lines = GraphvizGenealogyRenderer()._footnote_lines(genealogy, 760.0)

    assert lines == ("* Текст сноски",)


def test_marriage_descendant_stem_uses_exact_sign_center() -> None:
    renderer = GraphvizGenealogyRenderer()

    assert renderer._marriage_sign_xs(100.0, 160.0) == (124.0, 136.0)
    assert renderer._marriage_stem_x(124.0, 136.0) == 130.0
